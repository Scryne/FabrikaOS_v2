"""S14 birim testleri: saf rapor üreticisi (ReportLab PDF / openpyxl XLSX).

Üretici DB/IO'suz saf fonksiyondur: ``RaporIcerik`` → ``bytes``. PDF için sihirli bayt
(``%PDF``), XLSX için round-trip (openpyxl ile geri okuma) doğrulanır.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.modules.rapor.uretici import RaporIcerik, pdf_uret, xlsx_uret


def _ornek_icerik() -> RaporIcerik:
    return RaporIcerik(
        baslik="Üretim Raporu",
        alt_baslik="alfa • 2026-06-07 → 2026-07-07",
        kolonlar=["Gün", "Makine", "Toplam", "Sağlam", "OEE"],
        satirlar=[
            ["2026-07-01", "MAK-01", 480, 456, 0.87],
            ["2026-07-01", "MAK-02", 300, 300, 0.91],
            ["2026-07-02", "MAK-01", 500, 480, 0.85],
        ],
        ozet=["Toplam üretim: 1280", "Ortalama OEE: %87.7"],
    )


def test_pdf_uret_sihirli_bayt_ve_boyut() -> None:
    veri = pdf_uret(_ornek_icerik())
    assert isinstance(veri, bytes)
    assert veri.startswith(b"%PDF")
    assert len(veri) > 1000


def test_pdf_uret_bos_satirlar() -> None:
    icerik = RaporIcerik(
        baslik="Boş Rapor", alt_baslik="veri yok", kolonlar=["A", "B"], satirlar=[], ozet=[]
    )
    veri = pdf_uret(icerik)
    assert veri.startswith(b"%PDF")


def test_xlsx_uret_round_trip() -> None:
    icerik = _ornek_icerik()
    veri = xlsx_uret(icerik)
    wb = load_workbook(io.BytesIO(veri))
    ws = wb.active
    assert ws is not None
    hucreler = [[h.value for h in satir] for satir in ws.iter_rows()]
    duz = [deger for satir in hucreler for deger in satir]
    # Başlık + kolonlar + veri + özet dosyada bulunur.
    assert "Üretim Raporu" in duz
    assert "Gün" in duz and "OEE" in duz
    assert "MAK-01" in duz
    assert 480 in duz  # sayılar sayı olarak yazılır (metin değil)
    assert "Ortalama OEE: %87.7" in duz


def test_xlsx_uret_bos_satirlar() -> None:
    icerik = RaporIcerik(
        baslik="Boş Rapor", alt_baslik="veri yok", kolonlar=["A"], satirlar=[], ozet=[]
    )
    veri = xlsx_uret(icerik)
    wb = load_workbook(io.BytesIO(veri))
    assert wb.active is not None

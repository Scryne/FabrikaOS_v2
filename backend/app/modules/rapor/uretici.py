"""Saf rapor üreticisi — ``RaporIcerik`` → PDF (ReportLab) / XLSX (openpyxl) baytları.

DB/IO yoktur: çağıran katman (service) veriyi toplar, burada yalnız biçimlendirilir.
Fonksiyonlar CPU-bağımlıdır ve senkrondur; async uçlar ``asyncio.to_thread`` ile sarar.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


@dataclass(frozen=True, slots=True)
class RaporIcerik:
    """Biçimden bağımsız rapor içeriği: başlık + tablo + özet satırları."""

    baslik: str
    alt_baslik: str
    kolonlar: list[str]
    satirlar: list[list[Any]] = field(default_factory=list)
    ozet: list[str] = field(default_factory=list)


def _hucre_metni(deger: Any) -> str:
    if isinstance(deger, float):
        return f"{deger:.2f}"
    return "" if deger is None else str(deger)


def pdf_uret(icerik: RaporIcerik) -> bytes:
    """İçeriği tek tablolu, yatay A4 PDF'e döker."""
    tampon = io.BytesIO()
    belge = SimpleDocTemplate(
        tampon,
        pagesize=landscape(A4),
        title=icerik.baslik,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )
    stiller = getSampleStyleSheet()
    ince = ParagraphStyle("AltBaslik", parent=stiller["Normal"], textColor=colors.grey)

    akis: list[Any] = [
        Paragraph(icerik.baslik, stiller["Title"]),
        Paragraph(icerik.alt_baslik, ince),
        Spacer(1, 6 * mm),
    ]

    veri = [[_hucre_metni(k) for k in icerik.kolonlar]] + [
        [_hucre_metni(d) for d in satir] for satir in icerik.satirlar
    ]
    tablo = Table(veri, repeatRows=1)
    tablo.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9ca3af")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    akis.append(tablo)

    if icerik.ozet:
        akis.append(Spacer(1, 6 * mm))
        akis.extend(Paragraph(satir, stiller["Normal"]) for satir in icerik.ozet)

    belge.build(akis)
    return tampon.getvalue()


def xlsx_uret(icerik: RaporIcerik) -> bytes:
    """İçeriği tek sayfalı XLSX'e döker (sayılar sayı olarak yazılır)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # openpyxl yeni çalışma kitabında hep bir sayfa verir
    ws.title = "Rapor"

    kalin = Font(bold=True)
    ws.append([icerik.baslik])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.append([icerik.alt_baslik])
    ws.append([])

    ws.append(icerik.kolonlar)
    baslik_satiri = ws.max_row
    for sutun in range(1, len(icerik.kolonlar) + 1):
        ws.cell(row=baslik_satiri, column=sutun).font = kalin

    for satir in icerik.satirlar:
        ws.append(satir)

    if icerik.ozet:
        ws.append([])
        for ozet_satiri in icerik.ozet:
            ws.append([ozet_satiri])

    # Kolon genişlikleri: içerik uzunluğuna kaba uyum (okunabilirlik).
    for sutun, kolon in enumerate(icerik.kolonlar, start=1):
        genislik = max(
            len(str(kolon)),
            *(len(_hucre_metni(s[sutun - 1])) for s in icerik.satirlar if len(s) >= sutun),
            10,
        )
        ws.column_dimensions[get_column_letter(sutun)].width = min(genislik + 2, 40)

    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()
